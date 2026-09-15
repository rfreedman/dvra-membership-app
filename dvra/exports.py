"""CSV / XLSX / PDF exports (openpyxl + fpdf2)."""

from __future__ import annotations

import csv
import io
from datetime import datetime
from pathlib import Path
from typing import Any

# Unicode TTF so fpdf2 embeds text as UTF-16 (Identity-H), not a legacy 8-bit encoding.
_FONTS_DIR = Path(__file__).resolve().parent / "fonts"
_FONT_REGULAR = _FONTS_DIR / "DejaVuSans.ttf"
_FONT_BOLD = _FONTS_DIR / "DejaVuSans-Bold.ttf"
_PDF_FONT = "DvraSans"


def _cell(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    return str(value).strip()


def _new_landscape_pdf() -> Any:
    from fpdf import FPDF

    pdf = FPDF(orientation="L", unit="mm", format="Letter")
    pdf.set_margins(10, 10, 10)
    pdf.add_font(_PDF_FONT, "", str(_FONT_REGULAR))
    pdf.add_font(_PDF_FONT, "B", str(_FONT_BOLD))
    return pdf


def _csv_bytes(headers: list[str], rows: list[list[Any]]) -> bytes:
    buf = io.StringIO()
    buf.write("\ufeff")
    w = csv.writer(buf)
    w.writerow(headers)
    for row in rows:
        w.writerow([_cell(c) for c in row])
    return buf.getvalue().encode("utf-8")


def _xlsx_bytes(headers: list[str], rows: list[list[Any]], title: str) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    safe = title[:31] if title else "Report"
    for ch in "<>:\\/?*[]":
        safe = safe.replace(ch, "")
    ws.title = safe or "Report"
    for i, h in enumerate(headers, start=1):
        ws.cell(1, i, h)
    for r_i, row in enumerate(rows, start=2):
        for c_i, val in enumerate(row, start=1):
            ws.cell(r_i, c_i, _cell(val))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _pdf_landscape(heading: str, headers: list[str], widths: list[float], rows: list[list[str]]) -> bytes:
    pdf = _new_landscape_pdf()
    pdf.add_page()
    pdf.set_font(_PDF_FONT, "B", 11)
    pdf.cell(0, 8, heading)
    pdf.ln(10)
    pdf.set_font(_PDF_FONT, "B", 8)
    for i, h in enumerate(headers):
        pdf.cell(widths[i], 7, h, border=1)
    pdf.ln()
    pdf.set_font(_PDF_FONT, "", 7)
    for row in rows:
        for i, cell in enumerate(row):
            pdf.cell(widths[i], 6, str(cell), border=1)
        pdf.ln()
    return bytes(pdf.output())


def timestamp_stem(prefix: str, compact: bool = False) -> str:
    fmt = "%Y%m%d-%H%M%S" if compact else "%Y-%m-%d-%H%M%S"
    return f"{prefix}-{datetime.now().strftime(fmt)}"


# --- Members ---

MEMBER_KEYS = [
    "call_sign",
    "last_name",
    "first_name",
    "email",
    "phone",
    "address_street",
    "address_city",
    "address_state",
    "address_zip",
    "license_class",
    "membership_type",
    "arrl_member",
    "key_number",
    "paid_through",
]
MEMBER_HEADERS = [
    "Call sign",
    "Last name",
    "First name",
    "Email",
    "Phone",
    "Street",
    "City",
    "St",
    "ZIP",
    "License class",
    "Membership type",
    "ARRL",
    "Key #",
    "Paid through",
]


def _member_cell(key: str, value: Any) -> str:
    if key == "arrl_member":
        if isinstance(value, bool):
            return "yes" if value else "no"
        s = str(value).strip().lower()
        return "yes" if s in ("1", "true", "yes", "on") else "no"
    if key == "key_number":
        return str(int(value)) if value not in (None, "") else ""
    return _cell(value)


def _member_matrix(rows: list[dict]) -> list[list[str]]:
    return [[_member_cell(k, r.get(k)) for k in MEMBER_KEYS] for r in rows]


def members_csv(rows: list[dict]) -> bytes:
    return _csv_bytes(MEMBER_HEADERS, _member_matrix(rows))


def members_xlsx(rows: list[dict]) -> bytes:
    return _xlsx_bytes(MEMBER_HEADERS, _member_matrix(rows), "Members")


def members_pdf(rows: list[dict]) -> bytes:
    n = len(MEMBER_KEYS)
    w = [266 / n] * n
    pdf = _new_landscape_pdf()
    pdf.add_page()
    pdf.set_font(_PDF_FONT, "", 7)
    for i, h in enumerate(MEMBER_HEADERS):
        pdf.cell(w[i], 6, h, border=1)
    pdf.ln()
    for r in _member_matrix(rows):
        for i, cell in enumerate(r):
            text = cell if len(cell) <= 80 else cell[:77] + "..."
            pdf.cell(w[i], 5, text, border=1)
        pdf.ln()
    return bytes(pdf.output())


# --- Payments report ---

PAY_KEYS = [
    "member_id",
    "member_name",
    "call_sign",
    "payment_date",
    "paid_through",
    "membership_year",
    "membership_type",
    "form_number",
    "notes",
]
PAY_HEADERS = [
    "Member ID",
    "Member name",
    "Call sign",
    "Payment date",
    "Paid through",
    "Membership year",
    "Membership type",
    "Form #",
    "Notes",
]


def payments_csv(rows: list[dict]) -> bytes:
    return _csv_bytes(PAY_HEADERS, [[_cell(r.get(k)) for k in PAY_KEYS] for r in rows])


def payments_xlsx(rows: list[dict]) -> bytes:
    return _xlsx_bytes(PAY_HEADERS, [[_cell(r.get(k)) for k in PAY_KEYS] for r in rows], "Payment report")


def payments_pdf(rows: list[dict]) -> bytes:
    headers = ["Member", "Call", "Pay date", "Paid through", "Type", "Form #", "Notes"]
    widths = [50, 20, 24, 24, 30, 16, 100]
    data = []
    for r in rows:
        data.append(
            [
                _cell(r.get("member_name"))[:52],
                _cell(r.get("call_sign"))[:12],
                _cell(r.get("payment_date")),
                _cell(r.get("paid_through")),
                _cell(r.get("membership_type"))[:24],
                _cell(r.get("form_number"))[:12],
                _cell(r.get("notes"))[:400],
            ]
        )
    return _pdf_landscape("Payment report (filtered export)", headers, widths, data)


# --- Keyholders / rosters ---

def keyholders_csv(rows: list[dict]) -> bytes:
    headers = ["Member ID", "Key #", "Call sign", "Last name", "First name", "Email"]
    return _csv_bytes(headers, [[r["id"], r["key_number"], r["call_sign"], r["last_name"], r["first_name"], r["email"]] for r in rows])


def keyholders_xlsx(rows: list[dict]) -> bytes:
    headers = ["Member ID", "Key #", "Call sign", "Last name", "First name", "Email"]
    return _xlsx_bytes(
        headers,
        [[r["id"], r["key_number"], r["call_sign"], r["last_name"], r["first_name"], r["email"]] for r in rows],
        "Keyholders",
    )


def keyholders_pdf(rows: list[dict]) -> bytes:
    return _pdf_landscape(
        "Keyholders report",
        ["Member ID", "Key #", "Call", "Last name", "First name", "Email"],
        [22, 18, 28, 42, 42, 120],
        [
            [
                str(m["id"]),
                str(m["key_number"]),
                m["call_sign"][:14],
                m["last_name"][:40],
                m["first_name"][:40],
                m["email"][:80],
            ]
            for m in rows
        ],
    )


def roster_by_name_csv(rows: list[dict]) -> bytes:
    return _csv_bytes(["Last name", "First name", "Call sign"], [[r["last_name"], r["first_name"], r["call_sign"]] for r in rows])


def roster_by_name_xlsx(rows: list[dict]) -> bytes:
    return _xlsx_bytes(
        ["Last name", "First name", "Call sign"],
        [[r["last_name"], r["first_name"], r["call_sign"]] for r in rows],
        "Roster by name",
    )


def roster_by_name_pdf(rows: list[dict]) -> bytes:
    return _pdf_landscape(
        "Roster by name",
        ["Last name", "First name", "Call sign"],
        [70, 70, 130],
        [[m["last_name"][:48], m["first_name"][:48], m["call_sign"][:20]] for m in rows],
    )


def roster_by_callsign_csv(rows: list[dict]) -> bytes:
    return _csv_bytes(["Call sign", "Last name", "First name"], [[r["call_sign"], r["last_name"], r["first_name"]] for r in rows])


def roster_by_callsign_xlsx(rows: list[dict]) -> bytes:
    return _xlsx_bytes(
        ["Call sign", "Last name", "First name"],
        [[r["call_sign"], r["last_name"], r["first_name"]] for r in rows],
        "Roster by callsign",
    )


def roster_by_callsign_pdf(rows: list[dict]) -> bytes:
    return _pdf_landscape(
        "Roster by callsign",
        ["Call sign", "Last name", "First name"],
        [40, 70, 160],
        [[m["call_sign"][:14], m["last_name"][:48], m["first_name"][:48]] for m in rows],
    )
